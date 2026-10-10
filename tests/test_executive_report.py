from io import BytesIO
from datetime import datetime, timedelta
from uuid import uuid4

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models.department import Department
from app.models.meeting import Meeting
from app.routers.admin_analytics import get_executive_dataset
from tests.helpers import _auth_header, _create_room, _create_user, _make_token


def test_executive_report_excel_export_success(
    client: TestClient,
    db_session: Session,
) -> None:
    admin = _create_user(db_session, f"executive-excel-admin-{uuid4().hex}", role="admin")
    organizer = _create_user(db_session, f"executive-excel-user-{uuid4().hex}")
    department = Department(name=f"Executive Dept {uuid4().hex}")
    db_session.add(department)
    db_session.flush()
    organizer.department_id = department.id
    room = _create_room(db_session, f"Executive room {uuid4().hex}")
    start = datetime.now().replace(hour=10, minute=0, second=0, microsecond=0)
    meeting = Meeting(
        title="Executive report meeting",
        room_id=room.id,
        organizer_id=organizer.id,
        start_time=start,
        end_time=start + timedelta(hours=1),
        check_in_time=start,
        check_out_time=start + timedelta(minutes=45),
        status="COMPLETED",
    )
    db_session.add(meeting)
    db_session.commit()

    response = client.get(
        "/api/v1/admin/analytics/executive-report/export",
        params={"format": "excel", "department_id": department.id},
        headers=_auth_header(_make_token(admin)),
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert response.headers["content-disposition"].endswith(".xlsx\"")
    assert response.content.startswith(b"PK")
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    assert workbook.sheetnames == [
        "Tổng quan Executive",
        "Chi tiết Phòng & Thiết bị",
        "Lãng phí & No-Show",
    ]
    overview_values = list(workbook["Tổng quan Executive"].values)
    assert any("Tổng giờ họp thực tế" in str(row) for row in overview_values)


def test_executive_report_pdf_export_success(
    client: TestClient,
    db_session: Session,
) -> None:
    admin = _create_user(db_session, f"executive-pdf-admin-{uuid4().hex}", role="admin")
    response = client.get(
        "/api/v1/admin/analytics/executive-report/export",
        params={"format": "pdf"},
        headers=_auth_header(_make_token(admin)),
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/pdf")
    assert response.headers["content-disposition"].endswith(".pdf\"")
    assert response.content.startswith(b"%PDF")


def test_executive_report_requires_admin(
    client: TestClient,
    db_session: Session,
) -> None:
    employee = _create_user(db_session, f"executive-employee-{uuid4().hex}")
    response = client.get(
        "/api/v1/admin/analytics/executive-report/export",
        params={"format": "excel"},
        headers=_auth_header(_make_token(employee)),
    )

    assert response.status_code == 403


def test_executive_dataset_aggregates_hours_departments_and_no_show(
    db_session: Session,
) -> None:
    organizer = _create_user(db_session, f"executive-metrics-user-{uuid4().hex}")
    department = Department(name=f"Executive metrics dept {uuid4().hex}")
    db_session.add(department)
    db_session.flush()
    organizer.department_id = department.id
    room = _create_room(db_session, f"Executive metrics room {uuid4().hex}")
    now = datetime.combine(datetime.now().date(), datetime.min.time()) + timedelta(hours=12)
    completed = Meeting(
        title="Completed actual meeting",
        room_id=room.id,
        organizer_id=organizer.id,
        start_time=now - timedelta(hours=1),
        end_time=now + timedelta(hours=1),
        check_in_time=now - timedelta(minutes=45),
        check_out_time=now,
        status="COMPLETED",
    )
    no_show = Meeting(
        title="No-show booking",
        room_id=room.id,
        organizer_id=organizer.id,
        start_time=now,
        end_time=now + timedelta(minutes=30),
        status="CANCELLED_NO_SHOW",
    )
    db_session.add_all([completed, no_show])
    db_session.commit()

    data = get_executive_dataset(
        db_session,
        now.date(),
        now.date(),
        department_id=department.id,
    )

    assert data["summary"]["total_meetings"] == 2
    assert data["summary"]["total_actual_hours"] == 0.75
    assert data["summary"]["wasted_hours"] == 0.5
    assert data["summary"]["no_show_count"] == 1
    assert data["departments"] == [
        {
            "department_id": department.id,
            "department_name": department.name,
            "meeting_count": 2,
            "no_show_count": 1,
            "no_show_rate": 50.0,
        }
    ]
