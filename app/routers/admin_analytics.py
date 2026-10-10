"""Admin-only analytics endpoints."""

import io
import os
import unicodedata
from collections import Counter
from datetime import date, datetime, time, timedelta
from html import escape
from pathlib import Path
from typing import Any
from typing import Literal

from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.core.security import get_current_user
from app.db.session import get_db
from app.models.department import Department
from app.models.equipment import MeetingEquipment
from app.models.meeting import Meeting
from app.models.room import Room
from app.models.user import User

router = APIRouter()

EXCEL_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
EXECUTIVE_AVAILABLE_HOURS_PER_DAY = 8
EXECUTIVE_HOTSPOT_UTILIZATION = 0.8
EXECUTIVE_UNDERUTILIZED_UTILIZATION = 0.2


def get_current_admin_user(
    current_user: User = Depends(get_current_user),
) -> User:
    """Require an administrator account for analytics endpoints."""
    if (current_user.role or "").strip().casefold() != "admin":
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Chỉ quản trị viên mới được xem thống kê hủy cuộc họp.",
        )
    return current_user


@router.get("/cancellations")
def get_cancellation_analytics(
    start_date: date | None = Query(None),
    end_date: date | None = Query(None),
    room_id: int | None = Query(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_admin_user),
) -> dict[str, Any]:
    """Summarize manual cancellations and no-shows for admins."""
    del current_user

    today = date.today()
    selected_start = start_date or today.replace(day=1)
    selected_end = end_date or today
    if selected_start > selected_end:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="start_date không được lớn hơn end_date.",
        )

    range_start = datetime.combine(selected_start, time.min)
    # Use next midnight as an exclusive bound to include every time on end_date,
    # including fractional seconds after 23:59:59.
    range_end = datetime.combine(selected_end + timedelta(days=1), time.min)
    cancelled_at = func.coalesce(
        Meeting.cancelled_at,
        Meeting.updated_at,
        Meeting.created_at,
    )
    query = db.query(Meeting).filter(
        func.upper(Meeting.status).in_(["CANCELLED", "CANCELLED_NO_SHOW"]),
        cancelled_at >= range_start,
        cancelled_at < range_end,
    )
    if room_id is not None:
        query = query.filter(Meeting.room_id == room_id)
    meetings = query.all()

    cancelled_by_user = 0
    cancelled_no_show = 0
    total_wasted_hours = 0.0
    reasons: Counter[str] = Counter()
    rooms: Counter[tuple[int, str]] = Counter()
    offenders: Counter[tuple[int, str]] = Counter()

    for meeting in meetings:
        if (meeting.status or "").upper() == "CANCELLED_NO_SHOW":
            cancelled_no_show += 1
        else:
            cancelled_by_user += 1

        total_wasted_hours += max(
            0.0,
            (meeting.end_time - meeting.start_time).total_seconds() / 3600,
        )
        cancellation_reason = (meeting.cancellation_reason or "").strip()
        if not cancellation_reason:
            cancellation_reason = (
                "Tự động hủy (No-Show)"
                if (meeting.status or "").upper() == "CANCELLED_NO_SHOW"
                else "Không có lý do (Hủy cũ)"
            )
        reasons[cancellation_reason] += 1

        if meeting.room_id is not None:
            room_name = meeting.room.name if meeting.room else "Phòng đã xóa"
            rooms[(meeting.room_id, room_name)] += 1

        if meeting.organizer_id is not None:
            user_name = (
                meeting.organizer.full_name or meeting.organizer.username
                if meeting.organizer
                else "Người dùng đã xóa"
            )
            offenders[(meeting.organizer_id, user_name)] += 1

    return {
        "summary": {
            "total_cancelled": len(meetings),
            "cancelled_by_user": cancelled_by_user,
            "cancelled_no_show": cancelled_no_show,
            "total_wasted_hours": round(total_wasted_hours, 2),
        },
        "breakdown_by_reason": [
            {"reason": reason, "cancel_count": count}
            for reason, count in sorted(
                reasons.items(), key=lambda item: (-item[1], item[0].casefold())
            )
        ],
        "top_cancelled_rooms": [
            {"room_id": key[0], "room_name": key[1], "cancel_count": count}
            for key, count in sorted(
                rooms.items(), key=lambda item: (-item[1], item[0][1].casefold())
            )[:5]
        ],
        "top_offenders": [
            {"user_id": key[0], "user_name": key[1], "cancel_count": count}
            for key, count in sorted(
                offenders.items(), key=lambda item: (-item[1], item[0][1].casefold())
            )[:5]
        ],
    }


def generate_cancellation_excel(
    data: dict[str, Any],
    start_date: date,
    end_date: date,
) -> io.BytesIO:
    """Build an Excel workbook for the filtered cancellation analytics."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Thống kê Hủy phòng"
    sheet.merge_cells("A1:C1")
    sheet["A1"] = "RoomSync - Báo cáo Hủy phòng"
    sheet["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="2563EB")
    sheet["A1"].alignment = Alignment(horizontal="center")
    sheet.merge_cells("A2:C2")
    sheet["A2"] = f"Khoảng thời gian: {start_date:%d/%m/%Y} - {end_date:%d/%m/%Y}"
    sheet["A2"].alignment = Alignment(horizontal="center")

    summary = data.get("summary", {})
    total_cancelled = int(summary.get("total_cancelled", 0) or 0)
    no_show = int(summary.get("cancelled_no_show", 0) or 0)
    no_show_rate = no_show / total_cancelled if total_cancelled else 0
    rows: list[tuple[str, list[tuple[str, str]]]] = [
        (
            "Tổng quan",
            [
                ("Tổng số cuộc họp hủy", str(total_cancelled)),
                ("Số giờ lãng phí", f"{float(summary.get('total_wasted_hours', 0) or 0):.2f}"),
                ("Tỷ lệ no-show", f"{no_show_rate:.1%}"),
            ],
        ),
        (
            "Phòng có nhiều lịch hủy",
            [
                (str(item.get("room_name") or "Phòng đã xóa"), str(item.get("cancel_count", 0)))
                for item in data.get("top_cancelled_rooms", [])
            ],
        ),
        (
            "Lý do hủy phổ biến",
            [
                (str(item.get("reason") or "Không có lý do"), str(item.get("cancel_count", 0)))
                for item in data.get("breakdown_by_reason", [])
            ],
        ),
        (
            "Người đặt có lịch hủy nhiều",
            [
                (str(item.get("user_name") or "Người dùng đã xóa"), str(item.get("cancel_count", 0)))
                for item in data.get("top_offenders", [])
            ],
        ),
    ]

    current_row = 4
    section_fill = PatternFill("solid", fgColor="DBEAFE")
    header_fill = PatternFill("solid", fgColor="E2E8F0")
    for section_title, values in rows:
        sheet.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=2)
        section_cell = sheet.cell(row=current_row, column=1, value=section_title)
        section_cell.font = Font(bold=True, color="1E3A8A")
        section_cell.fill = section_fill
        current_row += 1

        headers = (
            ("Chỉ số", "Giá trị")
            if section_title == "Tổng quan"
            else (
                ("Phòng", "Số lần hủy")
                if section_title == "Phòng có nhiều lịch hủy"
                else (
                    ("Lý do", "Số lượng")
                    if section_title == "Lý do hủy phổ biến"
                    else ("Người đặt", "Số lần hủy/no-show")
                )
            )
        )
        for column, value in enumerate(headers, start=1):
            cell = sheet.cell(row=current_row, column=column, value=value)
            cell.font = Font(bold=True)
            cell.fill = header_fill
        current_row += 1

        for label, value in values:
            sheet.cell(row=current_row, column=1, value=label)
            sheet.cell(row=current_row, column=2, value=value)
            current_row += 1
        if not values:
            sheet.cell(row=current_row, column=1, value="Không có dữ liệu")
            sheet.cell(row=current_row, column=2, value="—")
            current_row += 1
        current_row += 1

    for column_index, column_cells in enumerate(sheet.columns, start=1):
        lengths = [
            len(str(cell.value or ""))
            for cell in column_cells
            if cell.__class__.__name__ != "MergedCell"
        ]
        max_length = max(lengths, default=12)
        column_letter = get_column_letter(column_index)
        sheet.column_dimensions[column_letter].width = min(max(max_length + 2, 14), 48)

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _pdf_font_name() -> tuple[str, bool]:
    """Use an installed Unicode font when available for Vietnamese PDF text."""
    candidates = [
        Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts" / "arial.ttf",
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"),
    ]
    for candidate in candidates:
        if candidate.is_file():
            font_name = "RoomSyncUnicode"
            if font_name not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(font_name, str(candidate)))
            return font_name, True
    return "Helvetica", False


def _pdf_text(value: Any, supports_unicode: bool) -> str:
    text = str(value if value is not None else "—")
    if not supports_unicode:
        text = "".join(
            char
            for char in unicodedata.normalize("NFKD", text)
            if not unicodedata.combining(char)
        )
    return escape(text)


def generate_cancellation_pdf(
    data: dict[str, Any],
    start_date: date,
    end_date: date,
) -> io.BytesIO:
    """Build a printable A4 PDF report for cancellation analytics."""
    font_name, supports_unicode = _pdf_font_name()
    output = io.BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title="RoomSync - Báo cáo Hủy phòng",
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "RoomSyncTitle",
        parent=styles["Title"],
        fontName=font_name,
        fontSize=18,
        leading=23,
        textColor=colors.HexColor("#1d4ed8"),
        alignment=TA_CENTER,
        spaceAfter=6,
    )
    body_style = ParagraphStyle(
        "RoomSyncBody",
        parent=styles["BodyText"],
        fontName=font_name,
        fontSize=9,
        leading=12,
    )
    heading_style = ParagraphStyle(
        "RoomSyncHeading",
        parent=styles["Heading2"],
        fontName=font_name,
        fontSize=12,
        leading=16,
        textColor=colors.HexColor("#1e3a8a"),
        spaceBefore=10,
        spaceAfter=5,
    )
    summary = data.get("summary", {})
    total_cancelled = int(summary.get("total_cancelled", 0) or 0)
    no_show = int(summary.get("cancelled_no_show", 0) or 0)
    no_show_rate = no_show / total_cancelled if total_cancelled else 0
    elements: list[Any] = [
        Paragraph(_pdf_text("RoomSync - Báo cáo Hủy phòng", supports_unicode), title_style),
        Paragraph(
            _pdf_text(
                f"Khoảng thời gian: {start_date:%d/%m/%Y} - {end_date:%d/%m/%Y}",
                supports_unicode,
            ),
            ParagraphStyle("RoomSyncDate", parent=body_style, alignment=TA_CENTER),
        ),
        Paragraph(
            _pdf_text(f"Ngày xuất báo cáo: {date.today():%d/%m/%Y}", supports_unicode),
            ParagraphStyle("RoomSyncExportDate", parent=body_style, alignment=TA_CENTER),
        ),
        Spacer(1, 8),
        Paragraph(_pdf_text("Tổng quan", supports_unicode), heading_style),
    ]

    def make_table(rows: list[list[Any]], widths: list[float]) -> Table:
        table = Table(rows, colWidths=widths, repeatRows=1, hAlign="LEFT")
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dbeafe")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#1e3a8a")),
                    ("FONTNAME", (0, 0), (-1, -1), font_name),
                    ("FONTNAME", (0, 0), (-1, 0), font_name),
                    ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                    ("LEADING", (0, 0), (-1, -1), 11),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 7),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                    ("TOPPADDING", (0, 0), (-1, -1), 6),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ]
            )
        )
        return table

    elements.append(
        make_table(
            [
                [
                    Paragraph(_pdf_text("Chỉ số", supports_unicode), body_style),
                    Paragraph(_pdf_text("Giá trị", supports_unicode), body_style),
                ],
                [
                    Paragraph(_pdf_text("Tổng số cuộc họp hủy", supports_unicode), body_style),
                    str(total_cancelled),
                ],
                [
                    Paragraph(_pdf_text("Số giờ lãng phí", supports_unicode), body_style),
                    f"{float(summary.get('total_wasted_hours', 0) or 0):.2f}",
                ],
                [
                    Paragraph(_pdf_text("Tỷ lệ no-show", supports_unicode), body_style),
                    f"{no_show_rate:.1%}",
                ],
            ],
            [105 * mm, 65 * mm],
        )
    )

    report_sections = [
        ("Phòng có nhiều lịch hủy", "top_cancelled_rooms", "room_name", "Phòng", "Số lần hủy"),
        ("Lý do hủy phổ biến", "breakdown_by_reason", "reason", "Lý do", "Số lượng"),
        ("Người đặt có lịch hủy nhiều", "top_offenders", "user_name", "Người đặt", "Số lần hủy/no-show"),
    ]
    for title, data_key, label_key, label_header, count_header in report_sections:
        elements.append(Paragraph(_pdf_text(title, supports_unicode), heading_style))
        values = data.get(data_key, [])
        rows = [
            [
                Paragraph(_pdf_text(label_header, supports_unicode), body_style),
                Paragraph(_pdf_text(count_header, supports_unicode), body_style),
            ]
        ]
        for item in values:
            rows.append(
                [
                    Paragraph(_pdf_text(item.get(label_key), supports_unicode), body_style),
                    str(item.get("cancel_count", 0)),
                ]
            )
        if not values:
            rows.append(
                [
                    Paragraph(_pdf_text("Không có dữ liệu", supports_unicode), body_style),
                    "—",
                ]
            )
        elements.append(make_table(rows, [105 * mm, 65 * mm]))

    document.build(elements)
    output.seek(0)
    return output


@router.get("/cancellations/export")
def export_cancellation_analytics(
    format: Literal["excel", "pdf"] = Query("excel"),
    start_date: date | None = Query(None),
    end_date: date | None = Query(None),
    room_id: int | None = Query(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_admin_user),
) -> Response:
    """Export the filtered cancellation analytics as Excel or PDF."""
    data = get_cancellation_analytics(
        start_date=start_date,
        end_date=end_date,
        room_id=room_id,
        db=db,
        current_user=current_user,
    )
    today = date.today()
    selected_start = start_date or today.replace(day=1)
    selected_end = end_date or today
    if format == "excel":
        content = generate_cancellation_excel(data, selected_start, selected_end)
        extension = "xlsx"
        media_type = EXCEL_MEDIA_TYPE
    else:
        content = generate_cancellation_pdf(data, selected_start, selected_end)
        extension = "pdf"
        media_type = "application/pdf"

    filename = f"Bao_Cao_Huy_Phong_{date.today():%Y%m%d}.{extension}"
    return Response(
        content=content.getvalue(),
        media_type=media_type,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _resolve_executive_period(
    start_date: date | None,
    end_date: date | None,
) -> tuple[date, date]:
    today = date.today()
    selected_start = start_date or today.replace(day=1)
    selected_end = end_date or today
    if selected_start > selected_end:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="start_date không được lớn hơn end_date.",
        )
    return selected_start, selected_end


def _calendar_days_inclusive(start_date: date, end_date: date) -> int:
    return (end_date - start_date).days + 1


def get_executive_dataset(
    db: Session,
    start_date: date,
    end_date: date,
    department_id: int | None = None,
) -> dict[str, Any]:
    """Aggregate meeting, room, department, equipment, and demand metrics."""
    range_start = datetime.combine(start_date, time.min)
    range_end = datetime.combine(end_date + timedelta(days=1), time.min)
    if department_id is not None and db.get(Department, department_id) is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Không tìm thấy phòng ban được yêu cầu.",
        )

    meetings_query = db.query(Meeting).filter(
        Meeting.start_time < range_end,
        Meeting.end_time > range_start,
    )
    if department_id is not None:
        meetings_query = meetings_query.join(
            User,
            Meeting.organizer_id == User.id,
        ).filter(User.department_id == department_id)
    meetings = meetings_query.all()

    cancelled_statuses = {"CANCELLED", "CANCELLED_NO_SHOW"}
    cancelled_meetings = [
        meeting for meeting in meetings
        if (meeting.status or "").upper() in cancelled_statuses
    ]
    active_meetings = [
        meeting for meeting in meetings
        if (meeting.status or "").upper() not in cancelled_statuses
    ]
    now = datetime.now()
    total_actual_hours = 0.0
    for meeting in active_meetings:
        if meeting.check_in_time is None:
            continue
        meeting_end = meeting.check_out_time or min(now, meeting.end_time)
        actual_start = max(meeting.check_in_time, range_start)
        actual_end = min(meeting_end, range_end)
        if actual_end > actual_start:
            total_actual_hours += (actual_end - actual_start).total_seconds() / 3600

    wasted_hours = 0.0
    for meeting in cancelled_meetings:
        wasted_start = max(meeting.start_time, range_start)
        wasted_end = min(meeting.end_time, range_end)
        if wasted_end > wasted_start:
            wasted_hours += (wasted_end - wasted_start).total_seconds() / 3600
    calendar_days = _calendar_days_inclusive(start_date, end_date)
    active_rooms = db.query(Room).filter(Room.is_active.is_(True)).order_by(Room.name).all()
    available_hours_per_room = calendar_days * EXECUTIVE_AVAILABLE_HOURS_PER_DAY
    total_available_hours = len(active_rooms) * available_hours_per_room
    utilization_rate = (
        total_actual_hours / total_available_hours
        if total_available_hours
        else 0.0
    )
    waste_rate = wasted_hours / total_available_hours if total_available_hours else 0.0

    equipment_totals: Counter[int] = Counter()
    equipment_query = (
        db.query(Meeting.room_id, func.sum(MeetingEquipment.quantity))
        .join(MeetingEquipment, MeetingEquipment.meeting_id == Meeting.id)
        .filter(
            Meeting.room_id.is_not(None),
            Meeting.start_time < range_end,
            Meeting.end_time > range_start,
            func.upper(Meeting.status).notin_(list(cancelled_statuses)),
        )
        .group_by(Meeting.room_id)
    )
    if department_id is not None:
        equipment_query = equipment_query.join(
            User,
            Meeting.organizer_id == User.id,
        ).filter(User.department_id == department_id)
    equipment_rows = equipment_query.all()
    for room_id, quantity in equipment_rows:
        equipment_totals[room_id] = int(quantity or 0)

    meetings_by_room: Counter[int] = Counter(
        meeting.room_id
        for meeting in active_meetings
        if meeting.room_id is not None
    )
    actual_hours_by_room: Counter[int] = Counter()
    for meeting in active_meetings:
        if meeting.room_id is None or meeting.check_in_time is None:
            continue
        meeting_end = meeting.check_out_time or min(now, meeting.end_time)
        actual_start = max(meeting.check_in_time, range_start)
        actual_end = min(meeting_end, range_end)
        if actual_end > actual_start:
            actual_hours_by_room[meeting.room_id] += (
                actual_end - actual_start
            ).total_seconds() / 3600

    room_metrics: list[dict[str, Any]] = []
    for room in active_rooms:
        available_hours = available_hours_per_room
        actual_hours = actual_hours_by_room[room.id]
        room_utilization = actual_hours / available_hours if available_hours else 0.0
        default_equipment = ", ".join(
            f"{allocation.equipment.name} × {allocation.quantity}"
            for allocation in room.default_equipments
            if allocation.equipment
        ) or "Chưa khai báo"
        if room_utilization >= EXECUTIVE_HOTSPOT_UTILIZATION:
            classification = "Quá tải"
        elif room_utilization < EXECUTIVE_UNDERUTILIZED_UTILIZATION:
            classification = "Sử dụng thấp"
        else:
            classification = "Cân bằng"
        room_metrics.append(
            {
                "room_id": room.id,
                "room_name": room.name,
                "meeting_count": meetings_by_room[room.id],
                "actual_hours": round(actual_hours, 2),
                "available_hours": available_hours,
                "utilization_rate": round(room_utilization * 100, 2),
                "classification": classification,
                "equipment_reserved_qty": equipment_totals[room.id],
                "default_equipment": default_equipment,
            }
        )
    room_metrics.sort(
        key=lambda item: (-item["meeting_count"], -item["utilization_rate"], item["room_name"].casefold())
    )

    department_stats: dict[int | None, dict[str, Any]] = {}
    for meeting in meetings:
        organizer = meeting.organizer
        department = organizer.department if organizer else None
        department_key = department.id if department else None
        metric = department_stats.setdefault(
            department_key,
            {
                "department_id": department_key,
                "department_name": department.name if department else "Chưa phân phòng ban",
                "meeting_count": 0,
                "no_show_count": 0,
            },
        )
        metric["meeting_count"] += 1
        if (meeting.status or "").upper() == "CANCELLED_NO_SHOW":
            metric["no_show_count"] += 1
    department_metrics = [
        {
            **metric,
            "no_show_rate": round(
                metric["no_show_count"] / metric["meeting_count"] * 100,
                2,
            ) if metric["meeting_count"] else 0.0,
        }
        for metric in department_stats.values()
    ]
    department_metrics.sort(
        key=lambda item: (-item["meeting_count"], item["department_name"].casefold())
    )

    weekday_names = (
        "Thứ Hai", "Thứ Ba", "Thứ Tư", "Thứ Năm", "Thứ Sáu", "Thứ Bảy", "Chủ Nhật"
    )
    peak_counts: Counter[tuple[int, int]] = Counter()
    for meeting in active_meetings:
        start_hour = meeting.start_time.hour
        slot_start = (start_hour // 2) * 2
        peak_counts[(meeting.start_time.weekday(), slot_start)] += 1
    peak_hours = [
        {
            "weekday": weekday_names[weekday],
            "weekday_index": weekday,
            "time_slot": f"{hour:02d}:00-{hour + 2:02d}:00",
            "meeting_count": count,
        }
        for (weekday, hour), count in sorted(
            peak_counts.items(),
            key=lambda item: (-item[1], item[0][0], item[0][1]),
        )
    ]

    cancellation_details = []
    for meeting in cancelled_meetings:
        organizer = meeting.organizer
        cancellation_details.append(
            {
                "meeting_id": meeting.id,
                "title": meeting.title,
                "status": (meeting.status or "").upper(),
                "room_name": meeting.room.name if meeting.room else "Phòng trực tuyến/đã xóa",
                "organizer_name": (
                    (organizer.full_name or organizer.username)
                    if organizer
                    else "Người dùng đã xóa"
                ),
                "department_name": (
                    organizer.department.name
                    if organizer and organizer.department
                    else "Chưa phân phòng ban"
                ),
                "start_time": meeting.start_time,
                "end_time": meeting.end_time,
                "wasted_hours": round(
                    max(
                        0.0,
                        (
                            min(meeting.end_time, range_end)
                            - max(meeting.start_time, range_start)
                        ).total_seconds() / 3600,
                    ),
                    2,
                ),
                "cancellation_reason": (
                    (meeting.cancellation_reason or "").strip()
                    or (
                        "Tự động hủy (No-Show)"
                        if (meeting.status or "").upper() == "CANCELLED_NO_SHOW"
                        else "Không có lý do"
                    )
                ),
            }
        )
    cancellation_details.sort(key=lambda item: item["start_time"], reverse=True)

    return {
        "period": {"start_date": start_date, "end_date": end_date},
        "summary": {
            "total_meetings": len(meetings),
            "active_meetings": len(active_meetings),
            "total_actual_hours": round(total_actual_hours, 2),
            "total_available_hours": total_available_hours,
            "utilization_rate": round(utilization_rate * 100, 2),
            "wasted_hours": round(wasted_hours, 2),
            "waste_rate": round(waste_rate * 100, 2),
            "cancelled_count": len(cancelled_meetings),
            "no_show_count": sum(
                1 for meeting in cancelled_meetings
                if (meeting.status or "").upper() == "CANCELLED_NO_SHOW"
            ),
        },
        "rooms": room_metrics,
        "departments": department_metrics,
        "peak_hours": peak_hours,
        "cancellations": cancellation_details,
        "hotspots": [
            item for item in room_metrics
            if item["classification"] == "Quá tải"
        ],
        "underutilized_rooms": [
            item for item in room_metrics
            if item["classification"] == "Sử dụng thấp"
        ],
    }


def generate_executive_excel(data: dict[str, Any], creator_name: str) -> io.BytesIO:
    """Create a styled, multi-sheet leadership report workbook."""
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Tổng quan Executive"
    rooms_sheet = workbook.create_sheet("Chi tiết Phòng & Thiết bị")
    waste_sheet = workbook.create_sheet("Lãng phí & No-Show")
    navy = "17365D"
    blue = "D9EAF7"
    light_gray = "F2F4F7"
    border = Border(
        left=Side(style="thin", color="D0D7DE"),
        right=Side(style="thin", color="D0D7DE"),
        top=Side(style="thin", color="D0D7DE"),
        bottom=Side(style="thin", color="D0D7DE"),
    )

    def style_sheet(sheet: Any) -> None:
        for row in sheet.iter_rows():
            for cell in row:
                cell.font = Font(name="Arial", size=10, color="1F2937")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.border = border
                if cell.row == 1:
                    cell.fill = PatternFill("solid", fgColor=navy)
                    cell.font = Font(name="Arial", size=15, bold=True, color="FFFFFF")
                elif cell.value and str(cell.value).startswith("## "):
                    cell.fill = PatternFill("solid", fgColor=blue)
                    cell.font = Font(name="Arial", size=11, bold=True, color=navy)
                    cell.value = str(cell.value)[3:]
                elif cell.row in (3,):
                    cell.fill = PatternFill("solid", fgColor=light_gray)
                    cell.font = Font(name="Arial", bold=True, color=navy)
        for column_index, cells in enumerate(sheet.columns, start=1):
            max_length = max((len(str(cell.value or "")) for cell in cells), default=10)
            sheet.column_dimensions[get_column_letter(column_index)].width = min(
                max(max_length + 2, 14), 46
            )
        sheet.freeze_panes = "A4"

    summary = data["summary"]
    period = data["period"]
    overview.append(["RoomSync | BÁO CÁO TỔNG HỢP HIỆU QUẢ SỬ DỤNG PHÒNG HỌP"])
    overview.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)
    overview.append(
        [
            f"Kỳ báo cáo: {period['start_date']:%d/%m/%Y} - {period['end_date']:%d/%m/%Y}",
            f"Người tạo: {creator_name}",
        ]
    )
    overview.append(["Công suất khả dụng ước tính: 8 giờ/phòng/ngày"])
    overview.append(["## Chỉ số tổng quan", "Giá trị", "Diễn giải", ""])
    kpis = [
        ("Tổng cuộc họp", summary["total_meetings"], "Trong kỳ báo cáo"),
        ("Tổng giờ họp thực tế", summary["total_actual_hours"], "Dựa trên check-in/check-out"),
        ("Tỷ lệ lấp đầy", f"{summary['utilization_rate']:.2f}%", "Giờ thực tế / giờ khả dụng"),
        ("Giờ lãng phí", summary["wasted_hours"], f"Tỷ lệ thất thoát {summary['waste_rate']:.2f}%"),
    ]
    for row in kpis:
        overview.append(list(row) + [""])
    overview.append(["## Xếp hạng phòng họp", "Số cuộc họp", "Giờ thực tế", "Lấp đầy"])
    for item in data["rooms"]:
        overview.append(
            [
                item["room_name"],
                item["meeting_count"],
                item["actual_hours"],
                f"{item['utilization_rate']:.2f}%",
            ]
        )
    overview.append(["## Phòng ban / Đội nhóm", "Số cuộc họp", "No-show", "Tỷ lệ no-show"])
    for item in data["departments"]:
        overview.append(
            [
                item["department_name"],
                item["meeting_count"],
                item["no_show_count"],
                f"{item['no_show_rate']:.2f}%",
            ]
        )
    overview.append(["## Khung giờ cao điểm", "Khung giờ", "Số cuộc họp", ""])
    for item in data["peak_hours"][:20]:
        overview.append(
            [item["weekday"], item["time_slot"], item["meeting_count"], ""]
        )

    rooms_sheet.append(["CHI TIẾT HIỆU SUẤT PHÒNG HỌP & THIẾT BỊ"])
    rooms_sheet.append(["Kỳ báo cáo", f"{period['start_date']:%d/%m/%Y} - {period['end_date']:%d/%m/%Y}"])
    rooms_sheet.append(
        [
            "Phòng họp",
            "Số cuộc họp",
            "Giờ thực tế",
            "Giờ khả dụng",
            "Tỷ lệ lấp đầy",
            "Phân loại",
            "Thiết bị đặt (SL)",
            "Thiết bị mặc định",
        ]
    )
    for item in data["rooms"]:
        rooms_sheet.append(
            [
                item["room_name"],
                item["meeting_count"],
                item["actual_hours"],
                item["available_hours"],
                f"{item['utilization_rate']:.2f}%",
                item["classification"],
                item["equipment_reserved_qty"],
                item["default_equipment"],
            ]
        )

    waste_sheet.append(["CHI TIẾT LÃNG PHÍ & NO-SHOW"])
    waste_sheet.append(["Kỳ báo cáo", f"{period['start_date']:%d/%m/%Y} - {period['end_date']:%d/%m/%Y}"])
    waste_sheet.append(
        [
            "Cuộc họp",
            "Trạng thái",
            "Phòng",
            "Người đặt",
            "Phòng ban",
            "Bắt đầu",
            "Kết thúc",
            "Giờ lãng phí",
            "Lý do",
        ]
    )
    for item in data["cancellations"]:
        waste_sheet.append(
            [
                item["title"],
                item["status"],
                item["room_name"],
                item["organizer_name"],
                item["department_name"],
                item["start_time"].strftime("%d/%m/%Y %H:%M"),
                item["end_time"].strftime("%d/%m/%Y %H:%M"),
                item["wasted_hours"],
                item["cancellation_reason"],
            ]
        )
    if not data["cancellations"]:
        waste_sheet.append(["Không có cuộc họp hủy/no-show trong kỳ"])

    for sheet in workbook.worksheets:
        style_sheet(sheet)
    for sheet in (overview,):
        for row in sheet.iter_rows(min_row=4, min_col=2):
            for cell in row:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def generate_executive_pdf(data: dict[str, Any], creator_name: str) -> io.BytesIO:
    """Create an A4 leadership report with summary, tables, and page numbering."""
    font_name, supports_unicode = _pdf_font_name()
    output = io.BytesIO()
    period = data["period"]
    summary = data["summary"]
    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=15 * mm,
        leftMargin=15 * mm,
        topMargin=20 * mm,
        bottomMargin=20 * mm,
        title="RoomSync - Báo cáo tổng hợp hiệu quả sử dụng phòng họp",
        author=creator_name,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ExecutiveTitle",
        parent=styles["Title"],
        fontName=font_name,
        fontSize=15,
        leading=19,
        textColor=colors.HexColor("#17365d"),
        alignment=TA_CENTER,
        spaceAfter=6,
    )
    body_style = ParagraphStyle(
        "ExecutiveBody",
        parent=styles["BodyText"],
        fontName=font_name,
        fontSize=8,
        leading=10,
    )
    section_style = ParagraphStyle(
        "ExecutiveSection",
        parent=styles["Heading2"],
        fontName=font_name,
        fontSize=11,
        leading=14,
        textColor=colors.HexColor("#17365d"),
        spaceBefore=8,
        spaceAfter=5,
    )
    small_style = ParagraphStyle(
        "ExecutiveSmall",
        parent=body_style,
        fontSize=7,
        leading=9,
    )

    def para(value: Any, style: ParagraphStyle = body_style) -> Paragraph:
        return Paragraph(_pdf_text(value, supports_unicode), style)

    def table(rows: list[list[Any]], widths: list[float]) -> Table:
        result = Table(rows, colWidths=widths, repeatRows=1, hAlign="LEFT")
        result.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#d9eaf7")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#17365d")),
                    ("FONTNAME", (0, 0), (-1, -1), font_name),
                    ("FONTSIZE", (0, 0), (-1, -1), 7.5),
                    ("GRID", (0, 0), (-1, -1), 0.45, colors.HexColor("#cbd5e1")),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 5),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ]
            )
        )
        return result

    kpis = [
        ("TỔNG CUỘC HỌP", str(summary["total_meetings"])),
        ("GIỜ HỌP THỰC TẾ", f"{summary['total_actual_hours']:.2f}"),
        ("TỶ LỆ LẤP ĐẦY", f"{summary['utilization_rate']:.2f}%"),
        ("GIỜ LÃNG PHÍ", f"{summary['wasted_hours']:.2f}"),
    ]
    kpi_rows = [
        [para(kpis[0][0]), para(kpis[1][0]), para(kpis[2][0]), para(kpis[3][0])],
        [para(kpis[0][1]), para(kpis[1][1]), para(kpis[2][1]), para(kpis[3][1])],
    ]
    elements: list[Any] = [
        para("ROOMSYNC", ParagraphStyle("ExecutiveLogo", parent=title_style, fontSize=11)),
        para("BÁO CÁO TỔNG HỢP HIỆU QUẢ SỬ DỤNG PHÒNG HỌP", title_style),
        para(
            f"Kỳ báo cáo: {period['start_date']:%d/%m/%Y} - {period['end_date']:%d/%m/%Y}"
            f" | Ngày xuất: {date.today():%d/%m/%Y} | Người tạo: {creator_name}",
            ParagraphStyle("ExecutiveMeta", parent=body_style, alignment=TA_CENTER),
        ),
        Spacer(1, 8),
        para("1. Tóm tắt dành cho Giám đốc", section_style),
    ]
    kpi_table = table(kpi_rows, [43 * mm] * 4)
    kpi_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#eef5fb")),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("TEXTCOLOR", (0, 1), (-1, 1), colors.HexColor("#17365d")),
                ("FONTSIZE", (0, 1), (-1, 1), 13),
            ]
        )
    )
    elements.append(kpi_table)
    elements.extend(
        [
            Spacer(1, 5),
            para(
                f"Giờ khả dụng ước tính: {summary['total_available_hours']} giờ; "
                "công suất quy ước 8 giờ/phòng/ngày; "
                f"tỷ lệ thất thoát do hủy/no-show: {summary['waste_rate']:.2f}% "
                f"({summary['no_show_count']} no-show trong {summary['cancelled_count']} lịch hủy).",
                body_style,
            ),
            para("2. Hiệu suất phòng họp và khung giờ cao điểm", section_style),
        ]
    )
    room_rows = [[
        para("Phòng", small_style),
        para("Số lịch", small_style),
        para("Giờ thực tế", small_style),
        para("Lấp đầy", small_style),
        para("Cảnh báo", small_style),
    ]]
    for item in data["rooms"][:12]:
        room_rows.append(
            [
                para(item["room_name"], small_style),
                str(item["meeting_count"]),
                f"{item['actual_hours']:.2f}",
                f"{item['utilization_rate']:.1f}%",
                para(item["classification"], small_style),
            ]
        )
    if len(room_rows) == 1:
        room_rows.append([para("Không có dữ liệu", small_style), "—", "—", "—", "—"])
    elements.append(table(room_rows, [57 * mm, 22 * mm, 28 * mm, 25 * mm, 38 * mm]))

    department_rows = [[
        para("Phòng ban/Đội nhóm", small_style),
        para("Số cuộc họp", small_style),
        para("No-show", small_style),
        para("Tỷ lệ", small_style),
    ]]
    for item in data["departments"][:8]:
        department_rows.append(
            [
                para(item["department_name"], small_style),
                str(item["meeting_count"]),
                str(item["no_show_count"]),
                f"{item['no_show_rate']:.1f}%",
            ]
        )
    if len(department_rows) == 1:
        department_rows.append([para("Không có dữ liệu", small_style), "—", "—", "—"])
    elements.extend(
        [
            Spacer(1, 7),
            table(department_rows, [80 * mm, 28 * mm, 24 * mm, 38 * mm]),
            Spacer(1, 7),
        ]
    )
    peak_rows = [[para("Thứ", small_style), para("Khung giờ", small_style), para("Số cuộc họp", small_style)]]
    for item in data["peak_hours"][:8]:
        peak_rows.append(
            [para(item["weekday"], small_style), item["time_slot"], str(item["meeting_count"])]
        )
    if len(peak_rows) == 1:
        peak_rows.append([para("Không có dữ liệu", small_style), "—", "—"])
    elements.append(table(peak_rows, [55 * mm, 55 * mm, 60 * mm]))

    no_show_rows = [[
        para("Cuộc họp", small_style),
        para("Phòng ban", small_style),
        para("Trạng thái", small_style),
        para("Giờ lãng phí", small_style),
    ]]
    for item in data["cancellations"][:12]:
        no_show_rows.append(
            [
                para(item["title"], small_style),
                para(item["department_name"], small_style),
                para(item["status"], small_style),
                f"{item['wasted_hours']:.2f}",
            ]
        )
    if len(no_show_rows) == 1:
        no_show_rows.append([para("Không có lịch hủy/no-show", small_style), "—", "—", "—"])
    no_show_rate = (
        summary["no_show_count"] / summary["total_meetings"] * 100
        if summary["total_meetings"]
        else 0.0
    )
    if no_show_rate >= 10:
        recommendation = (
            "Khuyến nghị: Tỷ lệ no-show từ 10% trở lên. Tăng nhắc lịch, yêu cầu "
            "check-in QR và xem xét tự động giải phóng phòng sau thời hạn check-in."
        )
    elif summary["utilization_rate"] >= 80:
        recommendation = (
            "Khuyến nghị: Công suất sử dụng cao. Cân nhắc mở rộng khung giờ/phòng "
            "khả dụng và phân bổ lại các phòng đang quá tải."
        )
    elif data["underutilized_rooms"]:
        recommendation = (
            f"Khuyến nghị: Có {len(data['underutilized_rooms'])} phòng sử dụng dưới 20%. "
            "Xem xét gom lịch, điều chỉnh quy mô phòng hoặc tái phân bổ thiết bị."
        )
    else:
        recommendation = (
            "Khuyến nghị: Duy trì theo dõi định kỳ công suất phòng và tỷ lệ no-show "
            "để cân bằng nguồn lực."
        )
    elements.extend(
        [
            para("3. Lãng phí và khuyến nghị quản trị", section_style),
            table(no_show_rows, [60 * mm, 42 * mm, 32 * mm, 36 * mm]),
            Spacer(1, 6),
            para(recommendation, body_style),
        ]
    )

    class NumberedCanvas(canvas.Canvas):
        def __init__(self, *args: Any, **kwargs: Any) -> None:
            super().__init__(*args, **kwargs)
            self._saved_page_states: list[dict[str, Any]] = []

        def showPage(self) -> None:
            self._saved_page_states.append(dict(self.__dict__))
            self._startPage()

        def save(self) -> None:
            total_pages = len(self._saved_page_states)
            for state in self._saved_page_states:
                self.__dict__.update(state)
                self.setFont(font_name, 8)
                self.setFillColor(colors.HexColor("#64748b"))
                self.drawString(15 * mm, 10 * mm, "RoomSync | Báo cáo quản trị")
                self.drawRightString(
                    A4[0] - 15 * mm,
                    10 * mm,
                    f"Trang {self._pageNumber} / {total_pages}",
                )
                canvas.Canvas.showPage(self)
            canvas.Canvas.save(self)

    document.build(elements, canvasmaker=NumberedCanvas)
    output.seek(0)
    return output


@router.get("/executive-report/export")
def export_executive_report(
    format: Literal["excel", "pdf"] = Query("excel"),
    start_date: date | None = Query(None),
    end_date: date | None = Query(None),
    department_id: int | None = Query(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_admin_user),
) -> Response:
    """Export the executive meeting utilization report."""
    selected_start, selected_end = _resolve_executive_period(start_date, end_date)
    data = get_executive_dataset(
        db,
        selected_start,
        selected_end,
        department_id=department_id,
    )
    creator_name = current_user.full_name or current_user.username
    if format == "excel":
        content = generate_executive_excel(data, creator_name)
        extension = "xlsx"
        media_type = EXCEL_MEDIA_TYPE
    else:
        content = generate_executive_pdf(data, creator_name)
        extension = "pdf"
        media_type = "application/pdf"
    filename = f"RoomSync_Executive_Report_{date.today():%Y%m%d}.{extension}"
    return Response(
        content=content.getvalue(),
        media_type=media_type,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
